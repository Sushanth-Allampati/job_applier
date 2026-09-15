"""
Extracts plain text from an uploaded resume file (PDF or DOCX).
"""
from __future__ import annotations

import io
from typing import Literal

import pdfplumber
import docx


def _extract_pdf_text(file_bytes: bytes) -> str:
    text_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_docx_text(file_bytes: bytes) -> str:
    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs]
    # Tables often hold skills/experience in resumes; pull those too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """
    Dispatches to the right extractor based on file extension.
    Raises ValueError for unsupported types.
    """
    lower = filename.lower()
    if lower.endswith(".pdf"):
        text = _extract_pdf_text(file_bytes)
    elif lower.endswith(".docx"):
        text = _extract_docx_text(file_bytes)
    else:
        raise ValueError("Unsupported file type. Please upload a .pdf or .docx resume.")

    if len(text.strip()) < 40:
        # Very little text usually means a scanned/image-based resume,
        # which most real ATS systems also fail to parse correctly.
        raise ValueError(
            "Could not extract meaningful text from this file. "
            "It may be a scanned image rather than a text-based document — "
            "this is also exactly the kind of resume that fails real ATS parsing."
        )
    return text
